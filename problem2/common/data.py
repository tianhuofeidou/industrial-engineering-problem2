"""只读加载原始数据；换模矩阵的行表示前件、列表示后件。"""
from dataclasses import dataclass
from pathlib import Path


def m62_duration(p: int) -> int:
    return (12 * p + 5) // 10


@dataclass(frozen=True)
class CaseData:
    case: int
    jobs: tuple[str, ...]
    processing: tuple[tuple[int, ...], ...]  # 工件 × 工位
    setup: tuple[tuple[int, ...], ...]  # 前件 × 后件，所有机器共用

    @property
    def n(self):
        return len(self.jobs)


def load_case(path: Path, case: int) -> CaseData:
    from openpyxl import load_workbook

    if case not in (1, 2):
        raise ValueError("case 必须是1或2")
    # 只加载，不调用 save；忽略工作表外围空白与不参与排产的对角元。
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        p_sheet, s_sheet = workbook["表2"], workbook["表4"]
        n, col, setup_row = (7, 3, 6) if case == 1 else (35, 12, 26)
        p_rows = list(p_sheet.iter_rows(min_row=6, max_row=16,
                                       min_col=col, max_col=col+n-1, values_only=True))
        machines = tuple(row[0] for row in p_sheet.iter_rows(min_row=7, max_row=16,
                                                           min_col=col-1, max_col=col-1,
                                                           values_only=True))
        if machines != tuple(f"M{k}" for k in range(1, 11)):
            raise ValueError("加工表的工位标签必须依次为M1至M10")
        s_rows = list(s_sheet.iter_rows(min_row=setup_row-1, max_row=setup_row+n-1,
                                       min_col=col-1, max_col=col+n-1, values_only=True))
        jobs = tuple(p_rows[0])
        expected = tuple(f"J{i+1}" for i in range(n))
        if jobs != expected or tuple(s_rows[0][1:]) != jobs or tuple(r[0] for r in s_rows[1:]) != jobs:
            raise ValueError("加工表与换模表的工件编号或顺序不一致")
        processing = tuple(tuple(p_rows[k+1][j] for k in range(10)) for j in range(n))
        for row in processing:
            if any(type(v) is not int or v <= 0 for v in row):
                raise ValueError("加工时间必须为正整数")
        setup = tuple(tuple(0 if i == j else s_rows[i+1][j+1] for j in range(n)) for i in range(n))
        if any(type(setup[i][j]) is not int or setup[i][j] < 0
               for i in range(n) for j in range(n) if i != j):
            raise ValueError("不同工件之间的换模时间必须为非负整数")
        return CaseData(case, jobs, processing, setup)
    finally:
        workbook.close()
